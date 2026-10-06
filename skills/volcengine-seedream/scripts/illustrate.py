#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Batch illustration for the seedream toolkit.

Sub-commands:
  gen      job.json [-o OUTDIR] [--result R.json] [--dry-run] [--jobs N]
           [--skip-existing] [--matte] [--bg '#RRGGBB'] [--feather-frac F]
  resolve  spec.json --result R.json -o spec.resolved.json
  embed    deck.pptx --placements P.json -o out.pptx

job.json (office -> seedream):
{
  "style": "flat vector illustration, unified blue-gray palette",
  "aspect_default": "16:9", "model": "flash", "size": "1.5K",
  "out_dir": "output/images",
  "items": [
    {"id": "s3_arch", "content": "...", "aspect": "16:9",
     "bg": "#0A0E27", "float": true}
  ]
}

result.json (seedream -> office):
{
  "items": [
    {"id","path","ok","model","size","aspect","unit_cost","error",
     "prep","seam_dev","skipped"}
  ],
  "total_cost": 0.24,
  "dry_run": false,
  "out_dir": "..."
}

resolve: copies a presentation spec (deck.py format) and fills each slide's
`image` from `image_id` using result.json, so the stock deck.py can embed it.

embed: inserts/replaces images in an existing pptx deck according to
placements.json (inches, 1-based slide indices). Operates on a copy.

--dry-run validates and quotes WITHOUT calling the API (no cost).
"""
from __future__ import annotations

import argparse
import concurrent.futures
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_image as gi  # reuse constants + pure helpers (no side effects on import)

GENERATOR = str(Path(__file__).resolve().parent / "generate_image.py")

DEFAULT_BG = "#0A0E27"
DEFAULT_FEATHER_FRAC = 0.06
DEFAULT_JOBS = 4
MAX_JOBS = 16


def compose_prompt(content: str, style: str | None) -> str:
    """Apply the global style lock to a per-item content string."""
    content = (content or "").strip().rstrip("。.")
    style = (style or "").strip()
    if style and style not in content:
        return f"{content}，{style}" if content else style
    return content


def sanitize_id(value: str) -> str:
    """Make an id safe for use as a file base name."""
    out = re.sub(r'[<>:"/\\|?*\s]+', "_", (value or "img").strip()).strip("_")
    return out or "img"


def item_plan(model: str, size: str, aspect: str | None) -> tuple[str | None, str | None, float | None]:
    """Return (family, tier, unit_cost). Any field may be None if unknown."""
    family = gi.normalize_model_name(model)
    tier = None
    try:
        width, height = gi.resolve_size(size, aspect)
        tier = gi.get_size_tier(width, height)
    except ValueError:
        tier = None
    cost = gi.PRICE_TABLE.get((family, tier)) if (family and tier) else None
    return family, tier, cost


# --------------------------------------------------------------------------- #
# Pure helpers for matte / placement (testable without PIL/numpy/pptx)
# --------------------------------------------------------------------------- #

def parse_bg(value: str) -> tuple[int, int, int]:
    """Parse '#RRGGBB' or 'RRGGBB' into (R, G, B)."""
    s = (value or "").strip()
    m = re.fullmatch(r"#?([0-9A-Fa-f]{6})", s)
    if not m:
        raise ValueError(f"invalid bg color '{value}', expected #RRGGBB")
    h = m.group(1)
    return tuple(int(h[i : i + 2], 16) for i in (0, 2, 4))


def border_mean_shift(border_mean: tuple, bg: tuple) -> tuple[int, int, int]:
    """Return per-channel shift needed to align border mean to bg."""
    return tuple(int(round(bg[i] - border_mean[i])) for i in range(3))


def apply_shift(pixel: tuple, shift: tuple) -> tuple[int, int, int]:
    """Add a per-channel shift and clamp to [0, 255]."""
    return tuple(max(0, min(255, int(round(pixel[i] + shift[i])))) for i in range(3))


def seam_deviation(border_mean_after: tuple, bg: tuple) -> int:
    """Max per-channel absolute deviation between a border mean and bg."""
    return max(abs(int(round(border_mean_after[i])) - bg[i]) for i in range(3))


def should_skip(path: Path) -> bool:
    """True if path exists and has non-zero size."""
    try:
        return path.is_file() and path.stat().st_size > 0
    except OSError:
        return False


_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")


def _existing_image(out_dir: Path, iid: str) -> Path | None:
    """Return the first existing non-empty image file for an id, if any."""
    for ext in _IMAGE_EXTS:
        candidate = out_dir / f"{iid}{ext}"
        if should_skip(candidate):
            return candidate
    return None


def _parse_saved_line(line: str) -> str | None:
    """Extract the file path from a generate_image.py 'SAVED:' output line."""
    stripped = line.strip()
    if not stripped.startswith("SAVED:"):
        return None
    return stripped.split("SAVED:", 1)[1].strip().rsplit("  (seed=", 1)[0]


def _compute_total_cost(items: list[dict]) -> tuple[float | None, str]:
    """Return (total_cost, display_text) for a list of result items.

    Only counts items that are ok and not skipped. If any such item has an
    unknown unit_cost, total_cost is None and text indicates unknown.
    """
    chargable = [it for it in items if it.get("ok") and not it.get("skipped")]
    has_unknown = any(it.get("unit_cost") is None for it in chargable)
    if has_unknown:
        return None, "未知（以控制台为准）"
    total = sum(it.get("unit_cost", 0.0) or 0.0 for it in chargable)
    return total, f"¥{total:.2f}"


def parse_placements(data: dict) -> list[dict]:
    """Validate and normalize a placements.json dict.

    Returns a list of normalized placements. Coordinates are in inches.
    """
    placements = data.get("placements") if isinstance(data, dict) else None
    if placements is None:
        raise ValueError("placements.json must contain a 'placements' list")
    if not isinstance(placements, list):
        raise ValueError("'placements' must be a list")

    out: list[dict] = []
    for idx, p in enumerate(placements):
        if not isinstance(p, dict):
            raise ValueError(f"placement {idx}: must be an object")

        slide = p.get("slide")
        if not isinstance(slide, int) or slide < 1:
            raise ValueError(f"placement {idx}: 'slide' must be a positive integer")

        image = p.get("image")
        if not image or not isinstance(image, str):
            raise ValueError(f"placement {idx}: 'image' path is required")

        def _num(key: str, default: float, positive: bool = False) -> float:
            v = p.get(key, default)
            try:
                v = float(v)
            except (TypeError, ValueError):
                raise ValueError(f"placement {idx}: '{key}' must be a number") from None
            if v < 0:
                raise ValueError(f"placement {idx}: '{key}' must be >= 0")
            if positive and v <= 0:
                raise ValueError(f"placement {idx}: '{key}' must be > 0")
            return v

        out.append(
            {
                "slide": slide,
                "image": image,
                "left": _num("left", 0.0),
                "top": _num("top", 0.0),
                "width": _num("width", 1.0, positive=True),
                "height": _num("height", 1.0, positive=True),
                "remove_pictures": bool(p.get("remove_pictures", False)),
            }
        )
    return out


# --------------------------------------------------------------------------- #
# Matte (requires Pillow + numpy; lazy imported)
# --------------------------------------------------------------------------- #

def _require_imaging() -> tuple:
    try:
        from PIL import Image
        import numpy as np

        return Image, np
    except Exception as e:
        raise RuntimeError(
            "Pillow and numpy are required for --matte. "
            "Use an environment with Pillow/numpy installed, e.g. "
            r"C:\ProgramData\anaconda3\envs\tools\python.exe"
        ) from e


def _border_band_mean(arr, np) -> tuple[float, float, float]:
    """Mean of the 1-pixel border ring, per RGB channel."""
    if arr.shape[0] < 2 or arr.shape[1] < 2:
        return (float(arr[..., 0].mean()), float(arr[..., 1].mean()), float(arr[..., 2].mean()))
    top = arr[0, :, :3]
    bottom = arr[-1, :, :3]
    left = arr[1:-1, 0, :3]
    right = arr[1:-1, -1, :3]
    border = np.concatenate([top.reshape(-1, 3), bottom.reshape(-1, 3), left.reshape(-1, 3), right.reshape(-1, 3)])
    return tuple(float(border[:, i].mean()) for i in range(3))


def _apply_shift_array(arr, shift: tuple[int, int, int], np):
    shifted = arr.astype(np.float32)
    for i in range(3):
        shifted[..., i] += shift[i]
    return np.clip(shifted, 0, 255).astype(np.uint8)


def _feather_alpha(h: int, w: int, ramp: float, np):
    y, x = np.indices((h, w))
    dist = np.minimum.reduce([x, y, w - 1 - x, h - 1 - y])
    return np.clip(dist / max(ramp, 1.0), 0, 1)


def _do_matte(src_path: str, dst_path: str, bg_hex: str, feather_frac: float, do_float: bool) -> int:
    """Produce a matted PNG and return its seam_deviation."""
    Image, np = _require_imaging()

    bg = parse_bg(bg_hex)
    img = Image.open(src_path).convert("RGBA")
    arr = np.array(img)
    rgb = arr[..., :3]

    border_before = _border_band_mean(rgb, np)
    shift = border_mean_shift(border_before, bg)
    shifted = _apply_shift_array(rgb, shift, np)

    if do_float:
        h, w = shifted.shape[:2]
        ramp = feather_frac * min(w, h)
        alpha = _feather_alpha(h, w, ramp, np)
        bg_arr = np.full_like(shifted, bg)
        composite = (
            shifted.astype(np.float32) * alpha[..., None]
            + bg_arr.astype(np.float32) * (1.0 - alpha[..., None])
        ).astype(np.uint8)
        out_arr = np.concatenate([composite, (alpha * 255).astype(np.uint8)[..., None]], axis=-1)
    else:
        out_arr = np.concatenate([shifted, np.full((*shifted.shape[:2], 1), 255, dtype=np.uint8)], axis=-1)

    border_after = _border_band_mean(out_arr[..., :3], np)
    seam_dev = seam_deviation(border_after, bg)

    Path(dst_path).parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(out_arr, "RGBA").save(dst_path, "PNG")
    return seam_dev


# --------------------------------------------------------------------------- #
# Generation
# --------------------------------------------------------------------------- #

def _build_item_record(it: dict, job: dict, index: int, results_len: int) -> dict:
    iid = sanitize_id(str(it.get("id") or f"img{results_len + 1}"))
    prompt = compose_prompt(it.get("content") or it.get("prompt") or "", job.get("style"))
    model = it.get("model") or job.get("model") or "flash"
    size = it.get("size") or job.get("size") or "1.5K"
    aspect = it.get("aspect") or job.get("aspect_default")
    family, tier, cost = item_plan(model, size, aspect)

    return {
        "id": iid,
        "prompt": prompt,
        "model": model,
        "size": size,
        "aspect": aspect,
        "family": family,
        "tier": tier,
        "unit_cost": cost,
        "bg": it.get("bg"),
        "float": bool(it.get("float", False)),
        "ok": True,
        "path": None,
        "error": None,
        "prep": None,
        "seam_dev": None,
        "skipped": False,
    }


def _run_one(item_rec: dict, out_dir: str, dry_run: bool, skip_existing: bool,
             matte: bool, bg: str, feather_frac: float) -> dict:
    """Generate one item (subprocess) and optionally matte it. Thread-safe."""
    iid = item_rec["id"]
    prompt = item_rec["prompt"]
    model = item_rec["model"]
    size = item_rec["size"]
    aspect = item_rec["aspect"]

    out_dir_path = Path(out_dir)
    out_dir_path.mkdir(parents=True, exist_ok=True)

    cmd = [
        sys.executable or "python",
        GENERATOR,
        "-p", prompt,
        "-m", model,
        "-s", size,
        "-f", iid,
        "-o", out_dir,
        "--no-preview",
    ]
    if aspect:
        cmd += ["-a", aspect]
    if dry_run:
        cmd += ["--dry-run"]

    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    rec = dict(item_rec)

    existing = _existing_image(out_dir_path, iid) if skip_existing else None

    if existing is not None:
        rec["skipped"] = True
        rec["path"] = str(existing.resolve())
    else:
        proc = subprocess.run(
            cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", env=env
        )
        out = (proc.stdout or "") + (proc.stderr or "")

        if proc.returncode != 0:
            rec["ok"] = False
            lines = [ln for ln in out.strip().splitlines() if ln.strip()]
            rec["error"] = (lines[-1] if lines else "unknown error")[:300]
        else:
            for ln in out.splitlines():
                parsed = _parse_saved_line(ln)
                if parsed:
                    rec["path"] = parsed
                    break

    # Optional matte post-processing.
    if matte and rec["ok"] and rec.get("path"):
        src = Path(rec["path"])
        prep_dir = out_dir_path / "_prep"
        prep_dir.mkdir(parents=True, exist_ok=True)
        prep_path = prep_dir / f"{iid}.png"

        item_bg = item_rec.get("bg") or bg
        do_float = bool(item_rec.get("float"))
        try:
            rec["seam_dev"] = _do_matte(str(src), str(prep_path), item_bg, feather_frac, do_float)
            rec["prep"] = str(prep_path.resolve())
        except Exception as e:
            rec["ok"] = False
            rec["error"] = f"matte failed: {e}"[:300]
            rec["prep"] = None

    status = "ok" if rec["ok"] else "FAIL"
    detail = rec.get("prep") or rec.get("path") or rec.get("error") or ""
    tag = "quote" if dry_run else ("skip" if rec.get("skipped") else "gen")
    print(f"[{tag}] {iid}: {status} {detail}")
    return rec


def run_gen(
    job_path: str,
    out_dir_override: str | None,
    result_path: str | None,
    dry_run: bool,
    jobs: int = DEFAULT_JOBS,
    skip_existing: bool = False,
    matte: bool = False,
    bg: str = DEFAULT_BG,
    feather_frac: float = DEFAULT_FEATHER_FRAC,
) -> int:
    job_path = str(Path(job_path).expanduser())
    try:
        job = json.loads(Path(job_path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"ERROR: invalid job JSON: {e}")
        return 2
    except OSError as e:
        print(f"ERROR: cannot read job file: {e}")
        return 2

    if not isinstance(job, dict):
        print("ERROR: job file must be a JSON object")
        return 2

    items = job.get("items") or []
    if not items:
        print("ERROR: job has no items")
        return 2
    if not isinstance(items, list):
        print("ERROR: job 'items' must be a list")
        return 2
    for idx, it in enumerate(items):
        if not isinstance(it, dict):
            print(f"ERROR: item {idx} is not an object")
            return 2

    job_dir = Path(job_path).resolve().parent
    out_dir = out_dir_override or job.get("out_dir") or str(job_dir / "images")
    result_path = result_path or str(Path(job_path).with_suffix("")) + ".result.json"

    if jobs > MAX_JOBS:
        print(f"WARNING: --jobs {jobs} exceeds max {MAX_JOBS}; clamped (Ark IPM ≈ 500/min, avoid flooding)")
        jobs = MAX_JOBS
    jobs = max(1, jobs)

    if feather_frac < 0 or feather_frac > 1:
        print(f"WARNING: --feather-frac {feather_frac} out of [0,1]; clamped")
        feather_frac = max(0.0, min(1.0, feather_frac))

    records = [_build_item_record(it, job, i, i) for i, it in enumerate(items)]

    seen_ids: set[str] = set()
    dupes: list[str] = []
    for rec in records:
        if rec["id"] in seen_ids:
            dupes.append(rec["id"])
        seen_ids.add(rec["id"])
    if dupes:
        print(f"ERROR: duplicate sanitized id(s): {dupes}")
        return 2

    results: list[dict] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=jobs) as ex:
        futures = {
            ex.submit(_run_one, rec, out_dir, dry_run, skip_existing, matte, bg, feather_frac): idx
            for idx, rec in enumerate(records)
        }
        for fut in concurrent.futures.as_completed(futures):
            idx = futures[fut]
            try:
                results.append((idx, fut.result()))
            except Exception as e:
                rec = dict(records[idx])
                rec["ok"] = False
                rec["error"] = f"internal error: {e}"[:300]
                print(f"[gen] {rec['id']}: FAIL {rec['error']}")
                results.append((idx, rec))

    results.sort(key=lambda t: t[0])
    ordered = [r for _, r in results]

    failed = sum(1 for r in ordered if not r["ok"])
    skipped = sum(1 for r in ordered if r.get("skipped"))

    total, cost_text = _compute_total_cost(ordered)

    res: dict = {
        "items": ordered,
        "total_cost": round(total, 2) if total is not None else None,
        "dry_run": dry_run,
        "out_dir": out_dir,
    }
    if matte:
        res["matte"] = {"bg": bg, "feather_frac": feather_frac}
    Path(result_path).write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"result: {result_path}")
    action = "预估费用" if dry_run else ("已出图" if not skipped else "已处理")
    print(
        f"共 {len(ordered)} 项，失败 {failed} 项，跳过 {skipped} 项；{action} ≈ {cost_text}"
    )
    return 0 if failed == 0 else 1


# --------------------------------------------------------------------------- #
# Resolve
# --------------------------------------------------------------------------- #

def run_resolve(spec_path: str, result_path: str, out_path: str) -> int:
    try:
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        res = json.loads(Path(result_path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"ERROR: invalid JSON: {e}")
        return 2
    except OSError as e:
        print(f"ERROR: cannot read input file: {e}")
        return 2

    img_map = {r["id"]: r["path"] for r in res.get("items", []) if r.get("ok") and r.get("path")}

    filled = 0
    for sl in spec.get("slides", []):
        iid = sl.get("image_id")
        if iid and not sl.get("image") and iid in img_map:
            sl["image"] = img_map[iid]
            filled += 1

    Path(out_path).write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"resolved {filled} image_id -> image; saved: {out_path}")

    missing = [sl.get("image_id") for sl in spec.get("slides", [])
               if sl.get("image_id") and not sl.get("image")]
    if missing:
        print(f"WARNING: 无对应结果图的 image_id: {missing}")
    return 0


# --------------------------------------------------------------------------- #
# Embed
# --------------------------------------------------------------------------- #

def _require_pptx() -> tuple:
    try:
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        return Presentation, Inches, MSO_SHAPE_TYPE
    except Exception as e:
        raise RuntimeError(
            "python-pptx is required for the embed command. "
            "Use an environment with python-pptx installed, e.g. "
            r"C:\ProgramData\anaconda3\envs\tools\python.exe"
        ) from e


def _shape_inventory(slide) -> list[dict]:
    out = []
    for shape in slide.shapes:
        if not hasattr(shape, "left"):
            continue
        try:
            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height
            if None in (left, top, width, height):
                continue
            out.append(
                {
                    "shape_id": shape.shape_id,
                    "shape_type": str(shape.shape_type),
                    "left_in": round(left / 914400.0, 4),
                    "top_in": round(top / 914400.0, 4),
                    "width_in": round(width / 914400.0, 4),
                    "height_in": round(height / 914400.0, 4),
                }
            )
        except Exception:
            continue
    return out


def run_embed(deck_path: str, placements_path: str, out_path: str) -> int:
    deck_path = Path(deck_path).expanduser()
    out_path = Path(out_path).expanduser()
    if not deck_path.is_file():
        print(f"ERROR: deck not found: {deck_path}")
        return 2
    if deck_path.resolve() == out_path.resolve():
        print("ERROR: output pptx must differ from input deck")
        return 2

    try:
        raw = json.loads(Path(placements_path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        print(f"ERROR: invalid placements JSON: {e}")
        return 2
    except OSError as e:
        print(f"ERROR: cannot read placements file: {e}")
        return 2

    placements = parse_placements(raw)
    if not placements:
        print("ERROR: no placements")
        return 2

    Presentation, Inches, MSO_SHAPE_TYPE = _require_pptx()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(str(deck_path), str(out_path))

    prs = Presentation(str(out_path))

    for p in placements:
        slide_idx = p["slide"] - 1
        if slide_idx >= len(prs.slides):
            print(f"ERROR: slide {p['slide']} does not exist")
            return 1
        slide = prs.slides[slide_idx]

        if p["remove_pictures"]:
            to_remove = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
            for s in to_remove:
                sp = s._element
                sp.getparent().remove(sp)

        img_path = Path(p["image"]).expanduser()
        if not img_path.is_file():
            print(f"ERROR: image not found: {img_path}")
            return 1
        slide.shapes.add_picture(
            str(img_path),
            Inches(p["left"]),
            Inches(p["top"]),
            Inches(p["width"]),
            Inches(p["height"]),
        )

    prs.save(str(out_path))

    print(f"saved: {out_path}")
    for idx, slide in enumerate(prs.slides, start=1):
        inv = _shape_inventory(slide)
        pictures = [s for s in inv if s["shape_type"] == "PICTURE (13)" or "PICTURE" in s["shape_type"]]
        print(f"slide {idx}: {len(pictures)} picture(s)")
        for pic in pictures:
            print(
                f"  picture L={pic['left_in']} T={pic['top_in']} "
                f"W={pic['width_in']} H={pic['height_in']}"
            )

    return 0


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def main() -> None:
    ap = argparse.ArgumentParser(description="Batch illustration (job.json -> images -> result.json) + deck resolve/embed.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    g = sub.add_parser("gen", help="generate images from a job.json")
    g.add_argument("job")
    g.add_argument("-o", "--out-dir", help="override output directory")
    g.add_argument("--result", help="result.json path (default: <job>.result.json)")
    g.add_argument("--dry-run", action="store_true", help="validate + quote only; no API call")
    g.add_argument("--jobs", type=int, default=DEFAULT_JOBS, help="concurrent generation workers (default 4)")
    g.add_argument("--skip-existing", action="store_true", help="skip items whose output image already exists")
    g.add_argument("--matte", action="store_true", help="normalize border color + optional feathered alpha")
    g.add_argument("--bg", default=DEFAULT_BG, help=f"matte target background color (default {DEFAULT_BG})")
    g.add_argument("--feather-frac", type=float, default=DEFAULT_FEATHER_FRAC, help="feather ramp as fraction of min(w,h) (default 0.06)")

    r = sub.add_parser("resolve", help="fill deck spec image_id fields from result.json")
    r.add_argument("spec")
    r.add_argument("--result", required=True)
    r.add_argument("-o", "--output", required=True)

    e = sub.add_parser("embed", help="insert/replace images in an existing pptx (operates on a copy)")
    e.add_argument("deck")
    e.add_argument("--placements", required=True, help="placements.json path")
    e.add_argument("-o", "--output", required=True, help="output pptx path")

    args = ap.parse_args()
    if args.cmd == "gen":
        sys.exit(
            run_gen(
                args.job,
                args.out_dir,
                args.result,
                args.dry_run,
                jobs=args.jobs,
                skip_existing=args.skip_existing,
                matte=args.matte,
                bg=args.bg,
                feather_frac=args.feather_frac,
            )
        )
    if args.cmd == "resolve":
        sys.exit(run_resolve(args.spec, args.result, args.output))
    sys.exit(run_embed(args.deck, args.placements, args.output))


if __name__ == "__main__":
    main()
